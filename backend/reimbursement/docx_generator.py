from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO

class ReimbursementDOCXGenerator:
    def __init__(self, reimbursement):
        self.reimbursement = reimbursement

    def generate(self):
        document = Document()
        
        # Header
        header = document.add_heading('EXPENSE REIMBURSEMENT VOUCHER', 0)
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Info Table
        table = document.add_table(rows=0, cols=2)
        
        def add_row(label, value):
            row_cells = table.add_row().cells
            row_cells[0].text = label
            row_cells[1].text = str(value)
            row_cells[0].paragraphs[0].runs[0].bold = True

        add_row('Voucher ID:', f"EXP-{self.reimbursement.id}")
        add_row('Date:', self.reimbursement.date)
        add_row('Employee Name:', self.reimbursement.employee.fullname)
        add_row('Employee ID (NIK):', self.reimbursement.employee.nik)
        add_row('Department:', self.reimbursement.employee.department.name if self.reimbursement.employee.department else '-')
        
        document.add_paragraph() # Spacer
        
        # Details
        document.add_heading('CLAIM DETAILS', level=1)
        
        details_table = document.add_table(rows=1, cols=2)
        details_table.style = 'Table Grid'
        hdr_cells = details_table.rows[0].cells
        hdr_cells[0].text = 'Category'
        hdr_cells[1].text = 'Description / Justification'
        
        row_cells = details_table.add_row().cells
        row_cells[0].text = self.reimbursement.category.name
        row_cells[1].text = self.reimbursement.description
        
        document.add_paragraph() # Spacer
        
        # Financials
        fin_table = document.add_table(rows=0, cols=2)
        
        def add_fin_row(label, amount):
            row = fin_table.add_row().cells
            row[0].text = label
            row[1].text = f"Rp {amount:,.0f}"
            row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

        add_fin_row('Requested Amount:', self.reimbursement.amount)
        add_fin_row('Approved Amount:', self.reimbursement.approved_amount or 0)
        
        document.add_paragraph()
        
        # Signatures
        document.add_paragraph('__________________________')
        document.add_paragraph('Employee Signature')
        
        document.add_paragraph()
        document.add_paragraph('__________________________')
        document.add_paragraph('Finance Approval')
        
        # Save to buffer
        f = BytesIO()
        document.save(f)
        return f.getvalue()
