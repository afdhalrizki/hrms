from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO

class PayslipDOCXGenerator:
    def generate(self, payslip):
        document = Document()
        
        # Header
        header = document.add_heading('PAYSLIP', 0)
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Company & Employee Info
        table = document.add_table(rows=4, cols=2)
        table.style = 'Table Grid'
        
        def set_cell(row, col, label, value):
            cells = table.rows[row].cells
            cells[col].text = f"{label}: {value}"

        set_cell(0, 0, "Employee", payslip.employee.fullname)
        set_cell(0, 1, "NIK", payslip.employee.nik)
        set_cell(1, 0, "Department", payslip.employee.department.name if payslip.employee.department else "-")
        set_cell(1, 1, "Role", payslip.employee.role.name if payslip.employee.role else "-")
        set_cell(2, 0, "Period", payslip.period.name)
        set_cell(2, 1, "Date", payslip.payment_date.strftime('%d %b %Y') if payslip.payment_date else "-")
        
        document.add_paragraph("\n")
        
        # Earnings & Deductions
        document.add_heading('Earnings', level=1)
        e_table = document.add_table(rows=3, cols=2)
        e_table.rows[0].cells[0].text = "Basic Salary"
        e_table.rows[0].cells[1].text = f"Rp {payslip.basic_salary:,.0f}"
        e_table.rows[1].cells[0].text = "Allowances"
        e_table.rows[1].cells[1].text = f"Rp {payslip.allowances:,.0f}"
        e_table.rows[2].cells[0].text = "Overtime"
        e_table.rows[2].cells[1].text = f"Rp {payslip.overtime_pay:,.0f}"
        
        document.add_paragraph("\n")
        
        document.add_heading('Deductions', level=1)
        d_table = document.add_table(rows=1, cols=2)
        d_table.rows[0].cells[0].text = "Deductions"
        d_table.rows[0].cells[1].text = f"Rp {payslip.deductions:,.0f}"
        
        document.add_paragraph("\n")
        
        # Summary
        summary = document.add_paragraph()
        summary.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = summary.add_run(f"NET PAY: Rp {payslip.net_pay:,.0f}")
        run.bold = True
        run.font.size = Pt(14)
        
        document.add_paragraph("\n" * 3)
        document.add_paragraph("Authorized Signature")
        
        f = BytesIO()
        document.save(f)
        return f.getvalue()
